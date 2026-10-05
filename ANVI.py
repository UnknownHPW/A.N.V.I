from datetime import datetime
import os
import yt_dlp
import psycopg2
import docx
from dotenv import load_dotenv
from docx.shared import Pt

load_dotenv("credentials.env")

print("olá william o que você quer fazer?")
resposta = input("")

if resposta == "baixar video":
    print("Digite o link do vídeo que deseja baixar:")
    link = input("")
    
    pasta_downloads = os.path.join(os.path.expanduser('~'), 'Downloads')
    caminho_saida = os.path.join(pasta_downloads, '%(title)s.%(ext)s')

    ydl_opts = {
        'outtmpl': caminho_saida,
    }

    print("Deseja baixar em mp3 ou mp4? (Digite 'mp3' ou 'mp4')")
    formato = input("")

    if formato == "mp3":
        ydl_opts.update({
            'format': 'bestaudio/best',
            'postprocessors': [{
                'key': 'FFmpegExtractAudio',
                'preferredcodec': 'mp3',
                'preferredquality': '192',
            }],
        })
    elif formato == "mp4":
        ydl_opts.update({
            'format': 'bv*[ext=mp4]+ba[ext=mp4]/b[ext=mp4] / best',
        })
    else:
        print("Formato não reconhecido, baixando no padrão melhor disponível.")

    with yt_dlp.YoutubeDL(ydl_opts) as ydl:
        ydl.download([link])
    
    print("Download concluído!")

elif resposta == "acessar o banco de dados":
    print("Conectando ao banco de dados...")
    try:
        conn = psycopg2.connect(
            host=os.getenv("DB_HOST"),
            database=os.getenv("DB_NAME"),
            user=os.getenv("DB_USER"),
            password=os.getenv("DB_PASSWORD"),
            port=os.getenv("DB_PORT")
        )
        cursor = conn.cursor()
        print("1. Adicionar um novo registro")
        print("2. Consultar registros")
        print("3. Atualizar um registro")
        print("4. Excluir um registro")

        opcao = input("Escolha uma opção: ")

        if opcao == "1":
            delete_query = "DELETE FROM compromissos WHERE data_compromisso < CURRENT_DATE"
            cursor.execute(delete_query)

            assunto = input("Digite o assunto do compromisso: ")
            data_usuario = input("Digite a data do compromisso (DD/MM/YYYY): ")
            
            data = datetime.strptime(data_usuario, "%d/%m/%Y").strftime("%Y-%m-%d")
            
            tipo = input("Digite o tipo do compromisso: ")
            descricao = input("Digite a descrição do compromisso: ")

            query = "INSERT INTO compromissos (assunto, data_compromisso, tipo, descricao) VALUES (%s, %s, %s, %s)"
            cursor.execute(query, (assunto, data, tipo, descricao))
            conn.commit()
            print("Registro adicionado com sucesso!")

        elif opcao == "2":
            mes = input("Digite o mês que deseja consultar (ou aperte Enter para ver TODOS): ")
            
            if mes == "":
                print("\n--- Todos os Compromissos ---")
                select_query = "SELECT * FROM compromissos ORDER BY data_compromisso ASC"
                cursor.execute(select_query)
            else:
                dia = input("Digite o dia (ou aperte Enter para ver o MÊS INTEIRO): ")
                
                if dia == "":
                    select_query = "SELECT * FROM compromissos WHERE EXTRACT(MONTH FROM data_compromisso) = %s ORDER BY data_compromisso ASC"
                    cursor.execute(select_query, (mes,))
                else:
                    select_query = "SELECT * FROM compromissos WHERE EXTRACT(MONTH FROM data_compromisso) = %s AND EXTRACT(DAY FROM data_compromisso) = %s ORDER BY data_compromisso ASC"
                    cursor.execute(select_query, (mes, dia))

            registros = cursor.fetchall()
            if not registros:
                print("Nenhum compromisso encontrado.")
            else:
                for registro in registros:
                    id_comp, assunto, data_comp, tipo, desc = registro
                    data_formatada = data_comp.strftime('%d/%m/%Y')
                    
                    print(f"\n--- Compromisso [{id_comp}] ---")
                    print(f"Assunto: {assunto}")
                    print(f"Data: {data_formatada}")
                    print(f"Tipo: {tipo}")
                    print(f"Descrição: {desc}")
                    print("-" * 30)

        elif opcao == "3":
            lembra = input("Digite o ID do compromisso que deseja atualizar (ou aperte Enter para ver todos): ")
            
            if lembra == "":
                print("\n--- Todos os Compromissos ---")
                select_query = "SELECT * FROM compromissos ORDER BY data_compromisso ASC"
                cursor.execute(select_query)

                registros = cursor.fetchall()
                if not registros:
                    print("Nenhum compromisso encontrado.")
                else:
                    for registro in registros:
                        id_comp, assunto, data_comp, tipo, desc = registro
                        data_formatada = data_comp.strftime('%d/%m/%Y')
                
                        print(f"\n--- Compromisso [{id_comp}] ---")
                        print(f"Assunto: {assunto}")
                        print(f"Data: {data_formatada}")
                        print(f"Tipo: {tipo}")
                        print(f"Descrição: {desc}")
                        print("-" * 30)
                    
                    lembra = input("Digite o ID do compromisso que deseja atualizar: ")

            if lembra != "":
                assunto = input("Digite o novo assunto do compromisso: ")
                data_usuario = input("Digite a nova data do compromisso (DD/MM/YYYY): ")
                data = datetime.strptime(data_usuario, "%d/%m/%Y").strftime("%Y-%m-%d")
                tipo = input("Digite o novo tipo do compromisso: ")
                descricao = input("Digite a nova descrição do compromisso: ")

                update_query = "UPDATE compromissos SET assunto = %s, data_compromisso = %s, tipo = %s, descricao = %s WHERE id_compromisso = %s"
                cursor.execute(update_query, (assunto, data, tipo, descricao, lembra))
                conn.commit()
                print("Registro atualizado com sucesso!")

        elif opcao == "4":
            lembra = input("Digite o ID do compromisso que deseja excluir (ou aperte Enter para ver todos): ")
            if lembra == "":
                print("\n--- Todos os Compromissos ---")
                select_query = "SELECT * FROM compromissos ORDER BY data_compromisso ASC"
                cursor.execute(select_query)

                registros = cursor.fetchall()
                if not registros:
                    print("Nenhum compromisso encontrado.")
                else:
                    for registro in registros:
                        id_comp, assunto, data_comp, tipo, desc = registro
                        data_formatada = data_comp.strftime('%d/%m/%Y')
                
                        print(f"\n--- Compromisso [{id_comp}] ---")
                        print(f"Assunto: {assunto}")
                        print(f"Data: {data_formatada}")
                        print(f"Tipo: {tipo}")
                        print(f"Descrição: {desc}")
                        print("-" * 30)
                    
                    lembra = input("Digite o ID do compromisso que deseja excluir: ")
            
            if lembra != "":
                delete_query = "DELETE FROM compromissos WHERE id_compromisso = %s"
                cursor.execute(delete_query, (lembra,))
                conn.commit()
                print("Registro excluído com sucesso!")
                    
            pass

        cursor.close()
        conn.close()
    except Exception as e:
        print(f"Erro ao conectar ao banco de dados: {e}")

elif resposta == "documento":
    print("Digite o arquivo que deseja alterar:")
    print("1. Declaração de frequencia")
    print("2. Justificativa de ponto")
    escolha = input("")
    
    if escolha == "1":
        caminho_docx = "decfreq.docx"
        if os.path.exists(caminho_docx):
            doc = docx.Document(caminho_docx)
            
            print("Digite o nome da criança:")
            novo_nome = input("")
            print("Digite a série da criança:")
            nova_serie = input("")
            print("Digite o turno:")
            novo_turno = input("")
            print("Digite a data:")
            nova_data = input("")

            dados = {
                '[nome]': novo_nome,
                '[serie]': nova_serie,
                '[turno]': novo_turno,
                '[data]': nova_data
            }

            for para in doc.paragraphs:
                for chave, valor in dados.items():
                    if chave in para.text:
                        para.text = para.text.replace(chave, valor)
                        for run in para.runs:
                            run.font.name = 'Arial'
                            run.font.size = Pt(12)
                            if chave == '[nome]' and novo_nome in run.text:
                                run.bold = True

            caminho_saida = os.path.join(os.path.expanduser('~'), 'Downloads', f'{novo_nome}.docx')
            doc.save(caminho_saida)
            print(f"Documento gerado com sucesso em: {caminho_saida}")
        else:
            print("O arquivo 'decfreq.docx' não foi encontrado na pasta do projeto.")

    elif escolha == "2":
        caminho_docx = "justponto.docx"
        if os.path.exists(caminho_docx):
            doc = docx.Document(caminho_docx)
            
            print("Digite o nome do servidor:")
            novo_nome = input("")
            print("Digite o dia:")
            nova_data = input("")
            print("Digite o horário (saída/entrada e hora):")
            novo_horario = input("")

            dados = {
                '[nome]': novo_nome,
                '[dia]': nova_data,
                '[horario]': novo_horario
            }

            for para in doc.paragraphs:
                for chave, valor in dados.items():
                    if chave in para.text:
                        para.text = para.text.replace(chave, valor)
                        for run in para.runs:
                            run.font.name = 'Arial'
                            run.font.size = Pt(12)
                            if chave == '[nome]' and novo_nome in run.text:
                                run.bold = True

            caminho_saida = os.path.join(os.path.expanduser('~'), 'Downloads', f'{novo_nome}_justificativa.docx')
            doc.save(caminho_saida)
            print(f"Documento gerado com sucesso em: {caminho_saida}")
        else:
            print("O arquivo 'justponto.docx' não foi encontrado na pasta do projeto.")
            
    else:
        print("Opção inválida! Escolha 1 ou 2.")